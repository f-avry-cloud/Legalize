#!/usr/bin/env python3
"""Génère les référentiels INPI (src/inpi/data/*.json) depuis les fichiers officiels.

Les tables de codes du Guichet unique ne sont pas exposées par une API : elles
sont publiées sous forme de classeurs, que ce script convertit en JSON pour que
l'application n'ait aucune table saisie à la main.

Fichiers attendus (téléchargeables sur inpi.fr, rubrique « Accès aux API
Guichet unique ») :
  --dictionnaire  inpi_dictionnaire-donnees-mandataire_<date>.xlsx
  --formes        inpi_liste-formes-juridiques-code-et-valideurs_<annee>.xlsx
  --activites     Categorisation_des_activites_<date>.xlsx   (facultatif)

Usage :
  pip install openpyxl
  python3 scripts/build-referentiels-inpi.py \
      --dictionnaire ~/inpi/dictionnaire.xlsx --formes ~/inpi/formes.xlsx
"""

import argparse
import json
import pathlib
import sys

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl requis : pip install openpyxl")

SORTIE = pathlib.Path(__file__).resolve().parent.parent / "src" / "inpi" / "data"

# Énumérations du dictionnaire reprises telles quelles (onglet → nom de sortie).
ENUMERATIONS = {
    "typeFormalite": "typeFormalite",
    "typePersonne": "typePersonne",
    "typeDePersonne": "typeDePersonne",
    "status": "statutFormalite",
    "statutPourLaFormalite": "statutPourLaFormalite",
    "rolePourEntreprise": "rolePourEntreprise",
    "diffusionINSEE": "diffusionINSEE",
    "formeExerciceActivitePrincipal": "formeExerciceActivitePrincipale",
    "succursaleOuFiliale": "succursaleOuFiliale",
    "modificationCapitalTypes": "modificationCapitalTypes",
    "typeAugmentationCapital": "typeAugmentationCapital",
    "typeReductionCapital": "typeReductionCapital",
    "typeDissolution": "typeDissolution",
    "motifCessation": "motifCessation",
    "evenementCessation": "evenementCessation",
    "typeVoie": "typeVoie",
    "documentExtension": "documentExtension",
}


def paires(ws):
    """Onglet à deux colonnes Code / Libellé → dict ordonné."""
    valeurs = {}
    for i, ligne in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not ligne or ligne[0] in (None, ""):
            continue
        valeurs[str(ligne[0]).strip()] = str(ligne[1] or "").strip()
    return valeurs


def ecrire(nom, donnees, source):
    chemin = SORTIE / f"{nom}.json"
    chemin.write_text(
        json.dumps({"_source": source, "valeurs": donnees}, ensure_ascii=False, indent=1) + "\n",
        encoding="utf-8",
    )
    taille = len(donnees) if isinstance(donnees, (dict, list)) else 1
    print(f"  {chemin.relative_to(SORTIE.parent.parent.parent)} — {taille} entrée(s)")


def dictionnaire(chemin):
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    source = pathlib.Path(chemin).name

    ecrire("roles", paires(wb["role"]), source)
    ecrire("pieces-justificatives", {
        str(l[0]).strip(): {"libelle": str(l[1] or "").strip(),
                            "nota": str(l[2]).strip() if len(l) > 2 and l[2] else None}
        for i, l in enumerate(wb["typeDocument"].iter_rows(values_only=True))
        if i and l and l[0]
    }, source)
    ecrire("evenements", paires(wb["events"]), source)

    enums = {}
    for onglet, nom in ENUMERATIONS.items():
        if onglet in wb.sheetnames:
            enums[nom] = paires(wb[onglet])
        else:
            print(f"  ! onglet absent : {onglet}", file=sys.stderr)
    ecrire("enumerations", enums, source)
    wb.close()


def formes_juridiques(chemin):
    """Onglet PM : le Guichet unique a son propre code, parfois ≠ du code INSEE."""
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    formes = {}
    for i, l in enumerate(wb["PM"].iter_rows(values_only=True)):
        if i < 2 or not l or not l[6]:
            continue
        code = str(l[6]).strip()
        libelle = (str(l[2]).strip() if l[2] else "") or (str(l[1]).strip() if l[1] else "")
        entree = {
            "libelle": libelle,
            "famille": str(l[1]).strip() if l[1] else "",
            "categorie": str(l[0]).strip() if l[0] else "",
            "codeInsee": str(l[7]).strip() if l[7] else code,
            "unipersonnelle": str(l[10]).strip().lower() == "oui" if l[10] else False,
            "creation": str(l[9]).strip().lower() == "oui" if l[9] else False,
        }
        # Une même forme apparaît en version pluripersonnelle et unipersonnelle :
        # on garde la pluripersonnelle comme libellé de référence.
        if code not in formes or not entree["unipersonnelle"]:
            formes[code] = entree
    ecrire("formes-juridiques", formes, pathlib.Path(chemin).name)
    wb.close()


def activites(chemin):
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    ws = wb["Nomenclature avec OF"]
    lignes = [l for l in ws.iter_rows(values_only=True)]
    entete = [str(c or "") for c in lignes[0]]
    try:
        i_code = next(i for i, c in enumerate(entete) if "code" in c.lower() and "ape" in c.lower())
    except StopIteration:
        i_code = 0
    categories = {}
    for l in lignes[1:]:
        if not l or not l[i_code]:
            continue
        categories[str(l[i_code]).strip()] = str(l[i_code + 1] or "").strip()
    ecrire("activites", categories, pathlib.Path(chemin).name)
    wb.close()


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--dictionnaire")
    p.add_argument("--formes")
    p.add_argument("--activites")
    args = p.parse_args()
    SORTIE.mkdir(parents=True, exist_ok=True)
    print("Génération des référentiels INPI :")
    if args.dictionnaire:
        dictionnaire(args.dictionnaire)
    if args.formes:
        formes_juridiques(args.formes)
    if args.activites:
        activites(args.activites)
    if not any([args.dictionnaire, args.formes, args.activites]):
        p.error("indiquer au moins un fichier source")
